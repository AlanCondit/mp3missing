"""Live Dylan album folders vs bundled studio track listings."""

from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from dylantrackcheck.catalog import load_catalog
from dylantrackcheck.compare import check_live_albums
from dylantrackcheck.names import (
    album_key,
    is_title_prefix,
    keys_match,
    parse_album_folder,
    track_key,
)
from dylantrackcheck.sheet import write_xlsx


def _touch(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"")


class NameTests(unittest.TestCase):
    def test_album_folder_pattern(self) -> None:
        self.assertEqual(
            parse_album_folder("1965 - Highway 61 Revisited"),
            (1965, "Highway 61 Revisited"),
        )
        self.assertIsNone(parse_album_folder("Highway 61 Revisited"))

    def test_album_key_strips_quotes_and_dashes(self) -> None:
        self.assertEqual(
            album_key('"Love and Theft"'),
            album_key("Love and Theft"),
        )
        self.assertEqual(
            album_key("Street‐Legal"),
            album_key("Street-Legal"),
        )

    def test_track_key_ignores_prefix_and_extension(self) -> None:
        self.assertEqual(
            track_key("01 - Like a Rolling Stone.mp3"),
            track_key("Like a Rolling Stone"),
        )
        self.assertEqual(
            track_key("Rainy Day Women #12 & 35.m4a"),
            track_key("Rainy Day Women ♯12 & 35"),
        )

    def test_track_key_treats_separators_alike_and_drops_punctuation(self) -> None:
        self.assertEqual(
            track_key("Love/Theft.mp3"),
            track_key("Love_Theft"),
        )
        self.assertEqual(
            track_key("Love-Theft"),
            track_key("Love Theft"),
        )
        self.assertEqual(
            track_key("01 - Don't Think Twice, It's All Right.mp3"),
            track_key("Dont Think Twice Its All Right"),
        )
        self.assertEqual(
            track_key("Mr. Tambourine Man"),
            track_key("Mr Tambourine Man"),
        )
        self.assertEqual(
            track_key("Just Like Tom Thumb’s Blues"),
            track_key("Just Like Tom Thumbs Blues"),
        )

    def test_track_key_equates_talkin_blowin_givin_with_ing(self) -> None:
        self.assertEqual(
            track_key("Talkin’ New York"),
            track_key("Talking New York"),
        )
        self.assertEqual(
            track_key("01 - Blowin' in the Wind.mp3"),
            track_key("Blowing in the Wind"),
        )
        self.assertEqual(
            track_key("Givin’ myself away"),
            track_key("Giving myself away"),
        )
        self.assertEqual(track_key("Talkin’ New York"), "talking new york")

    def test_track_key_strips_trailing_parentheticals(self) -> None:
        self.assertEqual(
            track_key("Girl From The North Country (With Johnny Cash).mp3"),
            track_key("Girl From the North Country"),
        )
        self.assertEqual(
            track_key("01 - Song (feat. Alice).mp3"),
            track_key("Song"),
        )
        self.assertEqual(
            track_key("Song (featuring Bob)"),
            track_key("Song"),
        )
        self.assertEqual(
            track_key("Song (Album Version)"),
            track_key("Song"),
        )
        self.assertEqual(
            track_key("Song (Remaster) (Official Audio)"),
            track_key("Song"),
        )
        # Display path still keeps the full name; only the key is cleaned.
        self.assertEqual(
            track_key("Girl From The North Country (With Johnny Cash)"),
            "girl from the north country",
        )

    def test_keys_match_trailing_now(self) -> None:
        self.assertTrue(
            keys_match(
                track_key("You're a Big Girl"),
                track_key("You're a Big Girl Now"),
            )
        )
        self.assertTrue(
            keys_match(track_key("Oh Sister"), track_key("Oh Sister Now.m4a"))
        )
        self.assertFalse(
            keys_match(track_key("Shelter"), track_key("Shelter From the Storm"))
        )

    def test_is_title_prefix_whole_word(self) -> None:
        self.assertTrue(
            is_title_prefix(
                track_key("Girl From the North Country"),
                track_key("Girl From the North Country Live"),
            )
        )
        self.assertFalse(
            is_title_prefix(track_key("It"), track_key("Idiot Wind"))
        )


class CheckTests(unittest.TestCase):
    def test_catalog_loads_forty_studio_albums(self) -> None:
        albums = load_catalog()
        self.assertEqual(len(albums), 40)
        highway = next(a for a in albums if a.title == "Highway 61 Revisited")
        self.assertEqual(highway.year, 1965)
        self.assertEqual(len(highway.tracks), 9)

    def test_only_present_albums_are_checked(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            dylan = root / "Bob Dylan"
            ok = dylan / "1965 - Highway 61 Revisited"
            bad = dylan / "1975 - Blood on the Tracks"
            for name in (
                "01 Like a Rolling Stone.mp3",
                "02 Tombstone Blues.mp3",
                "03 It Takes a Lot to Laugh, It Takes a Train to Cry.mp3",
                "04 From a Buick 6.mp3",
                "05 Ballad of a Thin Man.mp3",
                "06 Queen Jane Approximately.mp3",
                "07 Highway 61 Revisited.mp3",
                "08 Just Like Tom Thumb's Blues.mp3",
                "09 Desolation Row.mp3",
            ):
                _touch(ok / name)
            _touch(bad / "01 Tangled Up in Blue.mp3")
            _touch(bad / "99 Not On The Album.mp3")
            # Album not owned: Freewheelin' is absent — must not appear.
            report = check_live_albums(root)
            titles = {r.album_title for r in report.results}
            self.assertEqual(titles, {"Highway 61 Revisited", "Blood on the Tracks"})
            highway = next(r for r in report.results if r.album_title == "Highway 61 Revisited")
            blood = next(r for r in report.results if r.album_title == "Blood on the Tracks")
            self.assertTrue(highway.ok)
            self.assertFalse(blood.ok)
            self.assertGreaterEqual(blood.missing_count, 1)
            self.assertEqual(blood.extra_count, 1)

    def test_artist_folder_can_be_passed_directly(self) -> None:
        with TemporaryDirectory() as tmp:
            dylan = Path(tmp) / "Bob Dylan"
            album = dylan / "1969 - Nashville Skyline"
            for name in (
                "01 Girl from the North Country.mp3",
                "02 Nashville Skyline Rag.mp3",
                "03 To Be Alone with You.mp3",
                "04 I Threw It All Away.mp3",
                "05 Peggy Day.mp3",
                "06 Lay Lady Lay.mp3",
                "07 One More Night.mp3",
                "08 Tell Me That It Isn't True.mp3",
                "09 Country Pie.mp3",
                "10 Tonight I'll Be Staying Here with You.mp3",
            ):
                _touch(album / name)
            report = check_live_albums(dylan)
            self.assertEqual(report.albums_checked, 1)
            self.assertEqual(report.albums_ok, 1)

    def test_xlsx_summary_and_problems(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            dylan = root / "Bob Dylan"
            album = dylan / "1965 - Highway 61 Revisited"
            _touch(album / "01 Like a Rolling Stone.mp3")
            _touch(album / "Extra Song.mp3")
            report = check_live_albums(root)
            xlsx = root / "out.xlsx"
            write_xlsx(xlsx, report)
            book = load_workbook(xlsx)
            self.assertEqual(
                [c.value for c in book["Summary"][1]],
                [
                    "Year",
                    "Album",
                    "Folder",
                    "Expected tracks",
                    "Live tracks",
                    "Matched",
                    "Missing",
                    "Extra",
                    "Status",
                ],
            )
            self.assertEqual(book["Summary"]["I2"].value, "Mismatch")
            self.assertGreaterEqual(book["Problems"].max_row, 2)

    def test_trailing_now_matches_without_renaming_m4a(self) -> None:
        with TemporaryDirectory() as tmp:
            dylan = Path(tmp) / "Bob Dylan"
            album = dylan / "1975 - Blood on the Tracks"
            for name in (
                "01 Tangled Up in Blue.m4a",
                "02 Simple Twist of Fate.m4a",
                "03 You're a Big Girl.m4a",  # catalog has "... Now"
                "04 Idiot Wind.m4a",
                "05 You're Gonna Make Me Lonesome When You Go.m4a",
                "06 Meet Me in the Morning.m4a",
                "07 Lily, Rosemary and the Jack of Hearts.m4a",
                "08 If You See Her, Say Hello.m4a",
                "09 Shelter From the Storm.m4a",
                "10 Buckets of Rain.m4a",
            ):
                _touch(album / name)
            report = check_live_albums(dylan)
            result = report.results[0]
            self.assertTrue(result.ok)
            self.assertEqual(result.matched_count, 10)
            # File on disk is unchanged (matching only).
            self.assertTrue((album / "03 You're a Big Girl.m4a").is_file())

    def test_unique_prefix_on_same_album_matches(self) -> None:
        with TemporaryDirectory() as tmp:
            dylan = Path(tmp) / "Bob Dylan"
            album = dylan / "1969 - Nashville Skyline"
            for name in (
                "01 Girl From the North Country Something.m4a",
                "02 Nashville Skyline Rag.m4a",
                "03 To Be Alone with You.m4a",
                "04 I Threw It All Away.m4a",
                "05 Peggy Day.m4a",
                "06 Lay Lady Lay.m4a",
                "07 One More Night.m4a",
                "08 Tell Me That It Isn't True.m4a",
                "09 Country Pie.m4a",
                "10 Tonight I'll Be Staying Here with You.m4a",
            ):
                _touch(album / name)
            report = check_live_albums(dylan)
            result = report.results[0]
            self.assertTrue(result.ok)
            self.assertEqual(result.matched_count, 10)

    def test_ambiguous_prefix_does_not_match(self) -> None:
        with TemporaryDirectory() as tmp:
            dylan = Path(tmp) / "Bob Dylan"
            album = dylan / "1969 - Nashville Skyline"
            # Two live files both prefix-relate to the same expected title.
            _touch(album / "Girl From the North Country Live.m4a")
            _touch(album / "Girl From the North Country Demo.m4a")
            for name in (
                "02 Nashville Skyline Rag.m4a",
                "03 To Be Alone with You.m4a",
                "04 I Threw It All Away.m4a",
                "05 Peggy Day.m4a",
                "06 Lay Lady Lay.m4a",
                "07 One More Night.m4a",
                "08 Tell Me That It Isn't True.m4a",
                "09 Country Pie.m4a",
                "10 Tonight I'll Be Staying Here with You.m4a",
            ):
                _touch(album / name)
            report = check_live_albums(dylan)
            result = report.results[0]
            self.assertFalse(result.ok)
            self.assertEqual(result.missing_count, 1)
            self.assertEqual(result.extra_count, 2)


if __name__ == "__main__":
    unittest.main()
