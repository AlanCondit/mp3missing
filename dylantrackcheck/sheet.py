"""Write the Dylan track-check report as an Excel workbook."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from dylantrackcheck.compare import CheckReport


def write_xlsx(path: Path, report: CheckReport) -> None:
    """Save summary and per-track problem rows."""

    book = Workbook()

    summary = book.active
    summary.title = "Summary"
    summary.append(
        [
            "Year",
            "Album",
            "Folder",
            "Expected tracks",
            "Live tracks",
            "Matched",
            "Missing",
            "Extra",
            "Status",
        ]
    )
    for result in report.results:
        summary.append(
            [
                result.album_year,
                result.album_title,
                result.album_folder,
                result.expected_count,
                result.live_count,
                result.matched_count,
                result.missing_count,
                result.extra_count,
                "OK" if result.ok else "Mismatch",
            ]
        )
    for column, width in (
        ("A", 8),
        ("B", 34),
        ("C", 40),
        ("D", 16),
        ("E", 12),
        ("F", 10),
        ("G", 10),
        ("H", 10),
        ("I", 12),
    ):
        summary.column_dimensions[column].width = width

    problems = book.create_sheet("Problems")
    problems.append(
        ["Year", "Album", "Folder", "Issue", "Expected track", "Live file"]
    )
    for result in report.results:
        for issue in result.issues:
            problems.append(
                [
                    issue.album_year,
                    issue.album_title,
                    issue.album_folder,
                    issue.issue,
                    issue.expected_track,
                    issue.live_file,
                ]
            )
    for column, width in (
        ("A", 8),
        ("B", 34),
        ("C", 40),
        ("D", 18),
        ("E", 42),
        ("F", 42),
    ):
        problems.column_dimensions[column].width = width

    unmatched = book.create_sheet("Unmatched folders")
    unmatched.append(["Year", "Folder title", "Folder", "Live tracks"])
    for album in report.unmatched_live:
        unmatched.append(
            [album.year, album.title, album.path.name, len(album.tracks)]
        )
    for column, width in (("A", 8), ("B", 34), ("C", 40), ("D", 12)):
        unmatched.column_dimensions[column].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
